"""
Build or verify the shareable example workbook (``example.xlsx``).

The private v5 template (``2PP_print_protocol_v5.xlsx``, git-ignored) holds real
names, permIds and test rows. The example keeps only the header rows (layout,
colours, column widths, help comments) and may be edited by hand in Excel.

Two modes (run from printjob2openbis/)::

    python tools/make_example_workbook.py            # rebuild example.xlsx from the template
    python tools/make_example_workbook.py --check    # verify the (hand-edited) example.xlsx

Both scan the example for data from the template: every text of a removed data
row (operators, resins, furnaces, instruments, test rows) and every permId. A
rebuild that fails the scan is deleted; ``--check`` never modifies the file.
Run ``--check`` after every edit in Excel, before committing.
"""

import argparse
import re
import sys
import zipfile
from pathlib import Path
from typing import Dict, Set

import openpyxl
from openpyxl.workbook.workbook import Workbook

PACKAGE_DIR = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = PACKAGE_DIR / "2PP_print_protocol_v5.xlsx"
DEFAULT_EXAMPLE = PACKAGE_DIR / "example.xlsx"

#: Sheets of the example and the rows kept (the header rows). Other sheets are dropped.
HEADER_ROWS: Dict[str, int] = {
    "README": 1,
    "PrintJobs": 2,
    "Imaging": 2,
    "Lists": 1,
    "Columns": 1,
}

PERMID = re.compile(r"\d{17}-\d+")

#: Sheets whose data rows hold real data (names, permIds, test rows). Their removed
#: values are searched for in the example. README / Model / Columns rows are
#: documentation: their words also occur in the kept headers, so they are not searched.
DATA_SHEETS = {"PrintJobs", "Imaging", "Lists"}

#: Lists columns with real data. The other Lists columns are generic vocabulary
#: (Print status, Sample state, Yes/No …) that also appears in the kept help texts.
SENSITIVE_LIST_HEADERS = {
    "Operator", "Resin name", "Resin permId", "Furnace name", "Furnace permId",
    "Instrument name", "Instrument permId",
}


def strip_to_headers(workbook: Workbook) -> Set[str]:
    """
    Delete every row below the header rows (and the sheets not in the example), in place.

    Returns:
        Text values removed from the data sheets (the texts that must not leak).
    """
    removed: Set[str] = set()
    kept: Set[str] = set()
    for sheet in workbook.worksheets:
        keep = HEADER_ROWS.get(sheet.title, 0)
        list_headers = {c.column: c.value for c in sheet[1]} if sheet.title == "Lists" else {}
        for row in sheet.iter_rows():
            for cell in row:
                text = cell.value.strip() if isinstance(cell.value, str) else None
                vocabulary = sheet.title == "Lists" and \
                    list_headers.get(cell.column) not in SENSITIVE_LIST_HEADERS
                if text and (cell.row <= keep or vocabulary):
                    kept.add(text)
                elif text and sheet.title in DATA_SHEETS and not text.startswith("="):
                    removed.add(text)
                if cell.row > keep:
                    cell.comment = None
        for merged in list(sheet.merged_cells.ranges):
            if merged.max_row > keep:
                sheet.unmerge_cells(str(merged))
        if sheet.max_row > keep:
            sheet.delete_rows(keep + 1, sheet.max_row - keep)
    for title in [s for s in workbook.sheetnames if s not in HEADER_ROWS]:
        del workbook[title]
    # Short or shared words ("OK", "Yes", a header) are not evidence of a leak.
    return {text for text in removed - kept if len(text) >= 5}


def sensitive_texts(template: Path) -> Set[str]:
    """Texts of the template's data rows that must never appear in the example."""
    return strip_to_headers(openpyxl.load_workbook(template))


def leaked_texts(path: Path, sensitive: Set[str]) -> Set[str]:
    """Sensitive texts or permIds present anywhere in *path* (cells, comments, metadata)."""
    leaks: Set[str] = set()
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            xml = archive.read(name).decode("utf-8", errors="ignore")
            leaks.update(PERMID.findall(xml))
            leaks.update(text for text in sensitive if text in xml)
    return leaks


def build(template: Path, example: Path) -> int:
    """Rebuild *example* from *template*; delete it again if the scan finds data."""
    workbook = openpyxl.load_workbook(template)
    sensitive = strip_to_headers(workbook)
    workbook.save(example)
    leaks = leaked_texts(example, sensitive)
    if leaks:
        example.unlink()
        print(f"ERROR: {len(leaks)} value(s) from the template's data rows found; file deleted.")
        return 1
    print(f"Written {example} (header rows only, verified).")
    return 0


def check(template: Path, example: Path) -> int:
    """Verify a (hand-edited) *example* without changing it."""
    leaks = leaked_texts(example, sensitive_texts(template))
    if leaks:
        print(f"ERROR: {example.name} contains {len(leaks)} value(s) from the template's data "
              f"rows (names, permIds, test data). Remove them before committing.")
        return 1
    print(f"OK: {example.name} contains no data from the template.")
    return 0


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--check", action="store_true",
                        help="Verify the existing example instead of rebuilding it.")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--example", type=Path, default=DEFAULT_EXAMPLE)
    args = parser.parse_args()
    run = check if args.check else build
    sys.exit(run(args.template, args.example))
