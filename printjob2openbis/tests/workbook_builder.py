"""
Build synthetic v5 protocol workbooks for tests.

openpyxl cannot compute formulas, so formula columns (Print name, Resin permId,
Furnace permId …) are written as plain values. Only synthetic data here.
"""

from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

import openpyxl

from excel.column_mapping import (
    HEADER_ROW,
    REQUIRED_SHEETS,
    SHEET_COLUMNS,
    SHEET_IMAGING,
    SHEET_PRINTJOBS,
    Column,
)

FAKE_RESIN_PERMID = "20210101000000000-10001"
FAKE_SUBSTRATE_PERMID = "20210101000000000-12345"
FAKE_FURNACE_PERMID = "20210101000000000-30001"


def write_workbook(
    path: Path,
    prints: Sequence[Dict[str, Any]] = (),
    imaging: Sequence[Dict[str, Any]] = (),
    drop_headers: Optional[Dict[str, List[str]]] = None,
    sheets: Sequence[str] = tuple(REQUIRED_SHEETS),
) -> Path:
    """
    Write a v5-shaped workbook.

    Args:
        path: Target ``.xlsx`` file.
        prints: PrintJobs rows as field → value (see ``PRINTJOB_COLUMNS``).
        imaging: Imaging rows as field → value (see ``IMAGING_COLUMNS``).
        drop_headers: Sheet → headers to leave out (to test missing headers).
        sheets: Sheets to create.

    Returns:
        *path*.
    """
    drop_headers = drop_headers or {}
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    rows_per_sheet = {SHEET_PRINTJOBS: prints, SHEET_IMAGING: imaging}
    for name in sheets:
        sheet = workbook.create_sheet(name)
        if name not in SHEET_COLUMNS:
            continue
        columns: List[Column] = [
            c for c in SHEET_COLUMNS[name] if c.header not in drop_headers.get(name, [])
        ]
        previous_section = None
        for index, col in enumerate(columns, start=1):
            if col.section != previous_section:
                sheet.cell(HEADER_ROW - 1, index, col.section)
                previous_section = col.section
            sheet.cell(HEADER_ROW, index, col.header)
        for offset, values in enumerate(rows_per_sheet[name]):
            for index, col in enumerate(columns, start=1):
                if col.field in values:
                    sheet.cell(HEADER_ROW + 1 + offset, index, values[col.field])
    workbook.save(path)
    return path
