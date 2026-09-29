"""Tests for femtika_fill.mapping and femtika_fill.fill (synthetic workbook and run folders)."""

import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

_PROJECT_ROOT = Path(__file__).parent.parent
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))

import openpyxl

from excel.column_mapping import HEADER_ROW
from femtika_fill.fill import WorkbookLockedError, fill_workbook
from femtika_fill.mapping import FILL_MAP, number, on_off, range_value
from tests.test_femtika_reader import GCODE_BODY, write_run
from tests.workbook_builder import write_workbook

NOW = datetime(2026, 9, 30, 8, 15, 0)


def cells(path: Path, row: int = 3):
    """Header → value of one PrintJobs row."""
    ws = openpyxl.load_workbook(path)["PrintJobs"]
    headers = [c.value for c in ws[HEADER_ROW]]
    return {h: ws.cell(row, i).value for i, h in enumerate(headers, 1) if h}


class TestMapping(unittest.TestCase):

    def test_range_value(self):
        self.assertEqual(range_value((8.0, 8.0)), 8)
        self.assertEqual(range_value((2.5, 7.5)), "2.5–7.5")
        self.assertEqual(range_value((1000.0, 3000.0), 1 / 1000), "1–3")
        self.assertIsNone(range_value(None))

    def test_number_and_on_off(self):
        self.assertEqual(number(-1.6144135, 3), -1.614)
        self.assertEqual(number(10.0, 3), 10)
        self.assertEqual((on_off(True), on_off(False), on_off(None)), ("On", "Off", None))

    def test_manual_columns_never_mapped(self):
        fields = {t.field for t in FILL_MAP}
        self.assertFalse(fields & {"objective", "r", "structures_printed", "print_status"})


class TestFillWorkbook(unittest.TestCase):

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.logs = self.root / "logs"
        self.run = write_run(self.logs / "Test job_20260102_100000")
        self.excel = self.root / "protocol.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def _workbook(self, **row):
        values = {"print_code": "2PP-000001", "femtika_output_folder": self.run.name, **row}
        return write_workbook(self.excel, prints=[values])

    def test_fills_empty_cells(self):
        self._workbook()
        report = fill_workbook(self.excel, self.logs, now=NOW)
        row = cells(self.excel)
        self.assertEqual(row["Print date"], datetime(2026, 1, 2, 10, 0))
        self.assertEqual(row["Print duration [min]"], 30)
        self.assertEqual(row["Max laser power (calibration) [mW]"], 12)
        self.assertEqual(row["Laser power [mW]"], "2.5–7.5")
        self.assertEqual(row["Scan speed [mm/s]"], "1–3")
        self.assertEqual((row["Infinite FOV"], row["Tilt compensation"]), ("Off", "On"))
        self.assertEqual((row["Tilt alpha [°]"], row["Tilt beta [°]"]), (-1.5, 0.5))
        self.assertEqual((row["z start [µm]"], row["z end [µm]"]), (1000, 1010))
        self.assertEqual((row["Slicing distance [µm]"], row["Hatching distance [µm]"]), (0.25, 0.5))
        self.assertEqual((row["x start [µm]"], row["x end [µm]"]), (-100, 100))
        self.assertIsNone(row["Objective"])
        self.assertIsNone(row["Structures printed"])
        self.assertEqual(report.filled, len(FILL_MAP))
        self.assertTrue(report.saved)
        self.assertEqual(report.backup, self.root / "backups" / "protocol_20260930_081500.xlsx")
        self.assertTrue(report.backup.is_file())

    def test_filled_cells_kept(self):
        self._workbook(laser_power_mw=9, tilt_alpha_deg=-1.5, print_date=date(2026, 1, 2))
        report = fill_workbook(self.excel, self.logs, now=NOW)
        row = cells(self.excel)
        self.assertEqual(row["Laser power [mW]"], 9)
        [line] = report.rows[0].kept
        self.assertIn("Laser power [mW]: cell 9, run 2.5–7.5", line)
        self.assertEqual(sorted(report.rows[0].same), ["Print date", "Tilt alpha [°]"])

    def test_hyphen_range_counts_as_same(self):
        self._workbook(laser_power_mw="2.5-7.5")
        report = fill_workbook(self.excel, self.logs, now=NOW)
        self.assertIn("Laser power [mW]", report.rows[0].same)

    def test_dry_run_changes_nothing(self):
        self._workbook()
        before = self.excel.read_bytes()
        report = fill_workbook(self.excel, self.logs, dry_run=True)
        self.assertEqual(self.excel.read_bytes(), before)
        self.assertEqual(report.filled, len(FILL_MAP))
        self.assertFalse(report.saved)
        self.assertFalse((self.root / "backups").exists())

    def test_full_path_and_rows_without_folder(self):
        write_workbook(self.excel, prints=[
            {"print_code": "2PP-000001", "femtika_output_folder": str(self.run)},
            {"print_code": "2PP-000002"},
        ])
        report = fill_workbook(self.excel, None, now=NOW)
        self.assertEqual([r.row for r in report.rows], [3])
        self.assertIsNone(cells(self.excel, 4)["Print date"])

    def test_missing_folder_is_row_error(self):
        self._workbook(femtika_output_folder="nope")
        report = fill_workbook(self.excel, self.logs, now=NOW)
        self.assertIn("run folder not found", report.rows[0].error)
        self.assertFalse(report.saved)

    def test_missing_file_reported(self):
        (self.run / "calibration.json").unlink()
        self._workbook()
        report = fill_workbook(self.excel, self.logs, now=NOW)
        self.assertEqual(report.rows[0].missing,
                         ["Max laser power (calibration) [mW] (calibration.json not found)"])

    def test_gcode_job_leaves_slicing_empty_with_reason(self):
        write_run(self.run, body=GCODE_BODY)
        self._workbook()
        report = fill_workbook(self.excel, self.logs, now=NOW)
        self.assertIsNone(cells(self.excel)["Slicing distance [µm]"])
        self.assertIn("Slicing distance [µm] (job imports G-code (set in the slicer): fill by hand)",
                      report.rows[0].missing)

    def test_aborted_run_warns(self):
        write_run(self.run, aborted=True)
        self._workbook()
        report = fill_workbook(self.excel, self.logs, now=NOW)
        self.assertIn("aborted", report.rows[0].warnings[0])

    def test_open_in_excel_refused(self):
        self._workbook()
        (self.root / "~$protocol.xlsx").write_bytes(b"lock")
        with self.assertRaises(WorkbookLockedError):
            fill_workbook(self.excel, self.logs, now=NOW)

    def test_summary_lines(self):
        self._workbook(laser_power_mw=9)
        lines = fill_workbook(self.excel, self.logs, now=NOW).lines()
        self.assertTrue(lines[0].startswith(f"Row 3 (2PP-000001): {len(FILL_MAP) - 1}/{len(FILL_MAP)} "
                                            f"cells filled, 1 differ (kept)"))
        self.assertIn("Total: 1 row(s)", "\n".join(lines))


if __name__ == "__main__":
    unittest.main()
