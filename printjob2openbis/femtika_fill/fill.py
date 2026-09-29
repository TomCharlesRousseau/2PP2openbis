"""
Fill the PrintJobs sheet of the print protocol from the Femtika run folders.

For every row with a ``Femtika output folder``, the run folder is read and the
columns of :data:`~femtika_fill.mapping.FILL_MAP` are written, **only into
empty cells**. A filled cell whose value differs from the run is kept and
reported. Formula columns and the manual columns (Objective, R, Structures
printed) are never written.

The workbook is changed in place (its name stays, the file is shared) after a
timestamped backup copy ``backups/<name>_<YYYYMMDD_HHMMSS>.xlsx``. openpyxl
drops the cached formula values on save, so the file must be opened and saved
in Excel before ``check`` / ``upload``.
"""

import shutil
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl

from excel.column_mapping import FIRST_DATA_ROW, HEADER_ROW, SHEET_PRINTJOBS, column_by_field
from femtika_fill.mapping import FILL_MAP, Target
from femtika_fill.reader import FemtikaRun, read_run_folder, resolve_run_folder

BACKUP_DIR = "backups"


class WorkbookLockedError(Exception):
    """The workbook is open in Excel (or otherwise locked)."""


@dataclass
class RowReport:
    """What happened in one PrintJobs row."""

    row: int
    print_code: Any
    folder: Optional[Path] = None
    filled: List[str] = field(default_factory=list)  # headers written
    same: List[str] = field(default_factory=list)  # already filled with the same value
    kept: List[str] = field(default_factory=list)  # "header: cell value, run value"
    missing: List[str] = field(default_factory=list)  # "header (source)"
    warnings: List[str] = field(default_factory=list)
    error: Optional[str] = None

    def line(self, total: int) -> str:
        """One summary line, e.g. ``Row 3 (2PP-000001): 9/15 cells filled, …``."""
        head = f"Row {self.row} ({self.print_code})"
        if self.error:
            return f"{head}: ERROR {self.error}"
        parts = [f"{len(self.filled)}/{total} cells filled"]
        if self.same:
            parts.append(f"{len(self.same)} already filled")
        if self.kept:
            parts.append(f"{len(self.kept)} differ (kept)")
        if self.missing:
            parts.append(f"{len(self.missing)} missing")
        return f"{head}: " + ", ".join(parts)


@dataclass
class FillReport:
    """Result of :func:`fill_workbook`."""

    excel_path: Path
    rows: List[RowReport] = field(default_factory=list)
    backup: Optional[Path] = None
    saved: bool = False
    dry_run: bool = False

    @property
    def filled(self) -> int:
        return sum(len(r.filled) for r in self.rows)

    def lines(self) -> List[str]:
        """Per-row details and the total, ready to print."""
        total = len(FILL_MAP)
        out: List[str] = []
        for row in self.rows:
            out.append(row.line(total))
            out += [f"    filled: {h}" for h in row.filled]
            out += [f"    differs, kept: {k}" for k in row.kept]
            out += [f"    missing: {m}" for m in row.missing]
            out += [f"    WARNING: {w}" for w in row.warnings]
        errors = sum(1 for r in self.rows if r.error)
        verb = "would be filled" if self.dry_run else "filled"
        out.append("")
        out.append(f"Total: {len(self.rows)} row(s) with a Femtika output folder, "
                   f"{self.filled} cell(s) {verb}, "
                   f"{sum(len(r.kept) for r in self.rows)} differ (kept), "
                   f"{sum(len(r.missing) for r in self.rows)} missing, {errors} error(s).")
        if self.backup:
            out.append(f"Backup: {self.backup}")
        if self.saved:
            out.append("Saved. Open the file in Excel and save it once (formulas are recalculated) "
                       "before running check / upload.")
        elif self.dry_run:
            out.append("DRY-RUN: the workbook was not changed.")
        else:
            out.append("Nothing to fill: the workbook was not changed.")
        return out


def lock_file(excel_path: Path) -> Optional[Path]:
    """Excel's ``~$`` lock file of *excel_path* if it exists (the file is open in Excel)."""
    name = excel_path.name
    for candidate in ("~$" + name, "~$" + name[2:]):
        path = excel_path.with_name(candidate)
        if path.exists():
            return path
    return None


def _same(cell_value: Any, value: Any) -> bool:
    """True if a filled cell already holds *value* (numbers compared with tolerance, dates by day)."""
    if isinstance(value, datetime) and isinstance(cell_value, (date, datetime)):
        cell_day = cell_value.date() if isinstance(cell_value, datetime) else cell_value
        return cell_day == value.date()
    if isinstance(cell_value, (int, float)) and isinstance(value, (int, float)):
        return abs(cell_value - value) <= 1e-6 * max(1.0, abs(value))
    return str(cell_value).strip().replace("-", "–") == str(value).strip()


def _header_columns(ws) -> Dict[str, int]:
    """Header text (row 2) → column number."""
    return {str(c.value).strip(): c.column for c in ws[HEADER_ROW] if c.value is not None}


def _fill_row(ws, row: int, run: FemtikaRun, columns: Dict[str, int], report: RowReport,
              write: bool) -> None:
    """Fill one row from *run*; record everything in *report*."""
    if run.aborted:
        report.warnings.append("the run was aborted: check Print status.")
    for target in FILL_MAP:
        header = column_by_field(SHEET_PRINTJOBS, target.field).header
        value = target.value(run)
        if value is None:
            report.missing.append(f"{header} ({_missing_reason(run, target)})")
            continue
        cell = ws.cell(row, columns[header])
        if cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()):
            if write:
                cell.value = value
            report.filled.append(f"{header} = {_show(value)}")
        elif _same(cell.value, value):
            report.same.append(header)
        else:
            report.kept.append(f"{header}: cell {_show(cell.value)}, run {_show(value)}")


def _missing_reason(run: FemtikaRun, target: Target) -> str:
    """Why *target* has no value: an expected reason, the missing file, or its source."""
    if target.note and target.note in run.notes:
        return run.notes[target.note]
    source_file = target.source.split()[0]
    if source_file in run.missing_files:
        return f"{source_file} not found"
    related = [p for p in run.problems if p.startswith(source_file)]
    return related[0] if related else f"not in {target.source}"


def _show(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value)


def fill_workbook(excel_path: Path, logs_dir: Optional[Path] = None, dry_run: bool = False,
                  now: Optional[datetime] = None) -> FillReport:
    """
    Fill the empty cells of every PrintJobs row that has a Femtika output folder.

    Args:
        excel_path: The print protocol workbook (changed in place).
        logs_dir: Folder holding the run folders, for cells with a bare folder name.
        dry_run: Report what would be filled; change nothing.
        now: Timestamp of the backup name (default: now).

    Returns:
        The report (per row and total).

    Raises:
        FileNotFoundError: *excel_path* does not exist.
        WorkbookLockedError: The workbook is open in Excel.
        KeyError: A mapped column is missing in row 2 (layout changed).
    """
    excel_path = Path(excel_path)
    if not excel_path.is_file():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")
    lock = lock_file(excel_path)
    if lock is not None and not dry_run:
        raise WorkbookLockedError(
            f"{excel_path.name} is open in Excel ({lock.name} exists). Close it and run again.")

    workbook = openpyxl.load_workbook(excel_path)
    ws = workbook[SHEET_PRINTJOBS]
    columns = _header_columns(ws)
    needed = [column_by_field(SHEET_PRINTJOBS, f).header
              for f in ["print_code", "femtika_output_folder"] + [t.field for t in FILL_MAP]]
    absent = [h for h in needed if h not in columns]
    if absent:
        raise KeyError(f"Columns missing in {SHEET_PRINTJOBS} row {HEADER_ROW}: {', '.join(absent)}")

    report = FillReport(excel_path=excel_path, dry_run=dry_run)
    code_col = columns[column_by_field(SHEET_PRINTJOBS, "print_code").header]
    folder_col = columns[column_by_field(SHEET_PRINTJOBS, "femtika_output_folder").header]
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        cell = ws.cell(row, folder_col).value
        if cell is None or not str(cell).strip():
            continue
        row_report = RowReport(row=row, print_code=ws.cell(row, code_col).value)
        report.rows.append(row_report)
        folder = resolve_run_folder(str(cell), logs_dir)
        row_report.folder = folder
        if not folder.is_dir():
            row_report.error = f"run folder not found: {folder}"
            continue
        _fill_row(ws, row, read_run_folder(folder), columns, row_report, write=not dry_run)

    if dry_run or report.filled == 0:
        return report
    stamp = (now or datetime.now()).strftime("%Y%m%d_%H%M%S")
    backup = excel_path.parent / BACKUP_DIR / f"{excel_path.stem}_{stamp}{excel_path.suffix}"
    backup.parent.mkdir(exist_ok=True)
    shutil.copy2(excel_path, backup)
    report.backup = backup
    try:
        workbook.save(excel_path)
    except PermissionError as exc:
        raise WorkbookLockedError(
            f"{excel_path.name} cannot be written (open in Excel?). Backup kept: {backup}") from exc
    report.saved = True
    return report
