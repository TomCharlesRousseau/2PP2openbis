"""
Reader for the 2PP print protocol workbook (layout v5).

Reads the PrintJobs and Imaging sheets by header name into
:class:`~models.printjob.PrintJob` and :class:`~models.imaging.ImagingEvent`
objects. The reader does not validate content: structural problems (missing
sheets / headers, rows without key) are recorded in :class:`ProtocolData` for
the checker to report. Formula columns are read from Excel's cached values.
"""

from dataclasses import dataclass, field
from datetime import datetime, time
from pathlib import Path
from typing import Any, Dict, Iterator, List, Tuple

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from excel.column_mapping import (
    FIRST_DATA_ROW,
    HEADER_ROW,
    REQUIRED_SHEETS,
    SHEET_COLUMNS,
    SHEET_IMAGING,
    SHEET_KEY_FIELD,
    SHEET_PRINTJOBS,
    Column,
)
from models.cell import CellValue
from models.imaging import ImagingEvent
from models.printjob import PrintJob
from utils.logger import get_logger

logger = get_logger(__name__)

# Fields that do not make a row "data" on their own.
_NON_DATA_FIELDS = {"openbis_upload"}


@dataclass
class ProtocolData:
    """
    Everything read from the workbook.

    Attributes:
        file_path: The workbook that was read.
        prints: One entry per PrintJobs data row, in row order.
        imaging: One entry per Imaging data row, in row order.
        missing_sheets: Required sheets not found in the workbook.
        missing_headers: Sheet → expected headers not found in row 2.
        duplicate_headers: Sheet → headers found more than once in row 2 (first one is used).
        rows_without_key: Sheet → rows with input values but an empty key column
            (ignored as data; reported by the checker).
    """

    file_path: Path
    prints: List[PrintJob] = field(default_factory=list)
    imaging: List[ImagingEvent] = field(default_factory=list)
    missing_sheets: List[str] = field(default_factory=list)
    missing_headers: Dict[str, List[str]] = field(default_factory=dict)
    duplicate_headers: Dict[str, List[str]] = field(default_factory=dict)
    rows_without_key: Dict[str, List[int]] = field(default_factory=dict)


def normalise_header(value: Any) -> str:
    """Return a header with surrounding whitespace removed and inner whitespace collapsed."""
    return " ".join(str(value).split()) if value is not None else ""


def clean_cell(value: Any) -> CellValue:
    """
    Convert a raw openpyxl value to a model value.

    Text is stripped (empty text → ``None``); a datetime at midnight becomes a
    :class:`datetime.date`. Other values are returned unchanged.
    """
    if isinstance(value, str):
        text = value.strip()
        return text or None
    if isinstance(value, datetime):
        return value.date() if value.time() == time(0, 0) else value
    return value


def read_protocol(file_path: Path) -> ProtocolData:
    """
    Read the PrintJobs and Imaging sheets of the protocol workbook.

    Args:
        file_path: Path to the ``.xlsx`` file.

    Returns:
        :class:`ProtocolData` with the rows and any structural problems found.

    Raises:
        FileNotFoundError: If *file_path* does not exist.
    """
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"Excel file not found: {file_path}")

    logger.info(f"Reading Excel file: {file_path}")
    data = ProtocolData(file_path=file_path)
    workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    try:
        data.missing_sheets = [s for s in REQUIRED_SHEETS if s not in workbook.sheetnames]

        if SHEET_PRINTJOBS in workbook.sheetnames:
            rows = _read_sheet(workbook[SHEET_PRINTJOBS], SHEET_PRINTJOBS, data)
            data.prints = [PrintJob(row=r, **values) for r, values in rows]

        if SHEET_IMAGING in workbook.sheetnames:
            rows = _read_sheet(workbook[SHEET_IMAGING], SHEET_IMAGING, data)
            data.imaging = [ImagingEvent(row=r, **values) for r, values in rows]
    finally:
        workbook.close()

    logger.info(f"Read {len(data.prints)} print rows and {len(data.imaging)} imaging rows")
    return data


def _read_sheet(
    sheet: Worksheet, sheet_name: str, data: ProtocolData
) -> List[Tuple[int, Dict[str, CellValue]]]:
    """
    Read the data rows of one sheet as (row number, field → value).

    Missing / duplicate headers and rows without key are recorded in *data*.
    """
    columns = SHEET_COLUMNS[sheet_name]
    positions = _locate_columns(sheet, sheet_name, columns, data)
    key_field = SHEET_KEY_FIELD[sheet_name]
    input_fields = [
        c.field for c in columns if not c.formula and c.field not in _NON_DATA_FIELDS
    ]

    result: List[Tuple[int, Dict[str, CellValue]]] = []
    without_key: List[int] = []
    for row_number, cells in _iter_rows(sheet):
        values: Dict[str, CellValue] = {
            col.field: clean_cell(cells[pos]) if pos < len(cells) else None
            for col, pos in ((c, positions.get(c.field)) for c in columns)
            if pos is not None
        }
        has_input = any(values.get(f) is not None for f in input_fields)
        if key_field is None:
            if has_input:
                result.append((row_number, values))
        elif values.get(key_field) is not None:
            result.append((row_number, values))
        elif has_input:
            without_key.append(row_number)

    if without_key:
        data.rows_without_key[sheet_name] = without_key
    return result


def _locate_columns(
    sheet: Worksheet, sheet_name: str, columns: List[Column], data: ProtocolData
) -> Dict[str, int]:
    """
    Find the 0-based position of every expected column by its header in row 2.

    Returns:
        Field → column index for the headers found.
    """
    header_cells = next(
        sheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True), ()
    )
    found: Dict[str, int] = {}
    duplicates: List[str] = []
    for index, value in enumerate(header_cells):
        header = normalise_header(value)
        if not header:
            continue
        if header in found:
            duplicates.append(header)
        else:
            found[header] = index

    missing = [c.header for c in columns if c.header not in found]
    if missing:
        data.missing_headers[sheet_name] = missing
    if duplicates:
        data.duplicate_headers[sheet_name] = duplicates
    return {c.field: found[c.header] for c in columns if c.header in found}


def _iter_rows(sheet: Worksheet) -> Iterator[Tuple[int, Tuple[Any, ...]]]:
    """Yield (Excel row number, raw cell values) for every row from the first data row."""
    for offset, cells in enumerate(sheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True)):
        yield FIRST_DATA_ROW + offset, cells

